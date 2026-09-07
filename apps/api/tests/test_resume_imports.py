from __future__ import annotations

from collections.abc import Iterator
from io import BytesIO
from pathlib import Path

import pytest
from alembic import command
from alembic.config import Config
from docx import Document
from fastapi.testclient import TestClient

from jobpilot_api.api.security import JOBPILOT_EXTENSION_ORIGIN
from jobpilot_api.config import ApiSettings
from jobpilot_api.infrastructure.database.engine import sqlite_database_url
from jobpilot_api.main import create_app

WEB_ORIGIN = "http://127.0.0.1:5173"
WEB_HEADERS = {"Origin": WEB_ORIGIN, "Sec-Fetch-Site": "same-site"}
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MAX_SAFE_MULTIPART_BODY_BYTES = 10 * 1024 * 1024 + 64 * 1024


@pytest.fixture
def database_path(tmp_path: Path) -> Path:
    path = tmp_path / "jobpilot.db"
    config = Config("apps/api/alembic.ini")
    config.set_main_option("sqlalchemy.url", sqlite_database_url(path).render_as_string())
    command.upgrade(config, "head")
    return path


@pytest.fixture
def client(database_path: Path) -> Iterator[TestClient]:
    application = create_app(
        ApiSettings.from_environment({"JOBPILOT_CORS_ORIGINS": WEB_ORIGIN}),
        database_path=database_path,
    )
    with TestClient(application, base_url="http://127.0.0.1") as active_client:
        yield active_client


def test_parse_docx_returns_editable_preview_without_database_mutation(client: TestClient) -> None:
    before_resumes = client.get("/api/v1/resume-versions").json()
    before_profile = client.get("/api/v1/autofill-profile").json()

    response = _parse_docx(client, _resume_docx())

    assert response.status_code == 200, response.text
    preview = response.json()
    assert preview["fileType"] == "DOCX"
    assert "示例用户" in preview["extractedText"]
    assert preview["profileCandidates"]["personal"] == {
        "name": "示例用户",
        "phone": "13800138000",
        "email": "candidate@example.invalid",
        "currentCity": None,
    }
    assert preview["profileCandidates"]["education"][0]["school"] == "示例大学"
    assert preview["metrics"]["fileSizeBytes"] == len(_resume_docx())
    assert preview["metrics"]["extractedCharacterCount"] == len(preview["extractedText"])
    assert preview["metrics"]["pageCount"] is None
    assert isinstance(preview["metrics"]["parseLatencyMs"], int)
    assert client.get("/api/v1/resume-versions").json() == before_resumes
    assert client.get("/api/v1/autofill-profile").json() == before_profile


def test_parse_is_web_only_and_the_only_multipart_write_exception(client: TestClient) -> None:
    extension = client.post(
        "/api/v1/resume-imports/parse",
        headers={"Origin": JOBPILOT_EXTENSION_ORIGIN, "Sec-Fetch-Site": "none"},
        files={"file": ("resume.docx", _resume_docx(), DOCX_MIME)},
    )
    missing_origin = client.post(
        "/api/v1/resume-imports/parse",
        files={"file": ("resume.docx", _resume_docx(), DOCX_MIME)},
    )
    other_multipart = client.post(
        "/api/v1/resume-versions",
        headers=WEB_HEADERS,
        files={"file": ("resume.docx", _resume_docx(), DOCX_MIME)},
    )

    assert extension.status_code == 403
    assert extension.json()["error"]["code"] == "LOCAL_WRITE_FORBIDDEN"
    assert missing_origin.status_code == 403
    assert other_multipart.status_code == 415
    assert other_multipart.json()["error"]["code"] == "JSON_REQUIRED"


def test_parse_rejects_non_loopback_target(client: TestClient) -> None:
    response = client.post(
        "/api/v1/resume-imports/parse",
        headers={**WEB_HEADERS, "Host": "example.invalid"},
        files={"file": ("resume.docx", _resume_docx(), DOCX_MIME)},
    )

    assert response.status_code == 403
    assert response.json()["error"]["code"] == "LOCAL_WRITE_FORBIDDEN"


def test_parse_rejects_unexpected_multipart_parts(client: TestClient) -> None:
    response = client.post(
        "/api/v1/resume-imports/parse",
        headers=WEB_HEADERS,
        files=[
            ("file", ("resume.docx", _resume_docx(), DOCX_MIME)),
            ("notes", (None, "must not be accepted")),
        ],
    )

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "VALIDATION_ERROR"


def test_parse_returns_stable_error_for_malformed_docx(client: TestClient) -> None:
    response = _parse_docx(client, b"PK\x03\x04not-a-docx-package")

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "RESUME_DOCX_INVALID"


@pytest.mark.parametrize(
    ("filename", "mime", "content", "expected_status", "expected_code"),
    [
        ("resume.doc", "application/msword", b"not-doc", 415, "UNSUPPORTED_RESUME_FILE_TYPE"),
        (
            "resume.pdf",
            "application/pdf",
            b"MZ executable",
            422,
            "RESUME_FILE_SIGNATURE_MISMATCH",
        ),
        (
            "resume.docx",
            "application/pdf",
            b"PK\x03\x04wrong",
            422,
            "RESUME_FILE_SIGNATURE_MISMATCH",
        ),
        ("resume.pdf", "text/plain", b"%PDF-1.4 fake", 422, "RESUME_FILE_SIGNATURE_MISMATCH"),
    ],
)
def test_parse_rejects_unsupported_and_spoofed_files(
    client: TestClient,
    filename: str,
    mime: str,
    content: bytes,
    expected_status: int,
    expected_code: str,
) -> None:
    response = client.post(
        "/api/v1/resume-imports/parse",
        headers=WEB_HEADERS,
        files={"file": (filename, content, mime)},
    )

    assert response.status_code == expected_status
    assert response.json()["error"]["code"] == expected_code


def test_parse_rejects_more_than_ten_mib_before_document_parsing(client: TestClient) -> None:
    response = client.post(
        "/api/v1/resume-imports/parse",
        headers=WEB_HEADERS,
        files={"file": ("resume.pdf", b"%PDF-" + b"0" * (10 * 1024 * 1024), "application/pdf")},
    )

    assert response.status_code == 413
    assert response.json()["error"]["code"] == "RESUME_FILE_TOO_LARGE"


def test_parse_rejects_an_oversized_multipart_body_before_form_parsing(client: TestClient) -> None:
    response = client.post(
        "/api/v1/resume-imports/parse",
        headers={
            **WEB_HEADERS,
            "Content-Length": str(MAX_SAFE_MULTIPART_BODY_BYTES + 1),
        },
        files={"file": ("resume.docx", _resume_docx(), DOCX_MIME)},
    )

    assert response.status_code == 413
    assert response.json()["error"]["code"] == "RESUME_FILE_TOO_LARGE"


def test_parse_rejects_chunked_multipart_without_a_bounded_content_length(
    client: TestClient,
) -> None:
    boundary = "jobpilot-test-boundary"
    body = (
        (
            f"--{boundary}\r\n"
            'Content-Disposition: form-data; name="file"; filename="resume.docx"\r\n'
            f"Content-Type: {DOCX_MIME}\r\n\r\n"
        ).encode()
        + _resume_docx()
        + f"\r\n--{boundary}--\r\n".encode()
    )

    response = client.post(
        "/api/v1/resume-imports/parse",
        headers={**WEB_HEADERS, "Content-Type": f"multipart/form-data; boundary={boundary}"},
        content=iter((body,)),
    )

    assert response.status_code == 413
    assert response.json()["error"]["code"] == "RESUME_FILE_TOO_LARGE"


def test_parse_does_not_log_untrusted_filename_or_resume_text(
    client: TestClient, caplog: pytest.LogCaptureFixture
) -> None:
    private_filename = "private-name-13800138000.docx"
    private_text = "private-resume-marker"

    response = client.post(
        "/api/v1/resume-imports/parse",
        headers=WEB_HEADERS,
        files={"file": (private_filename, _resume_docx(private_text), DOCX_MIME)},
    )

    assert response.status_code == 200, response.text
    log_text = caplog.text
    assert private_filename not in log_text
    assert private_text not in log_text


def test_confirm_can_create_resume_only_and_profile_only_without_touching_other_data(
    client: TestClient,
) -> None:
    resume_only = client.post(
        "/api/v1/resume-imports/confirm",
        headers=WEB_HEADERS,
        json={"resumeVersion": {"name": "导入简历 2026-09-07", "content": "编辑后的正文"}},
    )

    assert resume_only.status_code == 200, resume_only.text
    assert resume_only.json()["resumeVersion"]["name"] == "导入简历 2026-09-07"
    assert resume_only.json()["profile"] is None
    assert client.get("/api/v1/autofill-profile").json() == {"profile": None}

    profile_only = client.post(
        "/api/v1/resume-imports/confirm",
        headers=WEB_HEADERS,
        json={
            "profileImport": {
                "personal": {"name": "示例用户", "email": "candidate@example.invalid"},
                "education": [{"school": "示例大学", "major": "信息管理"}],
            }
        },
    )

    assert profile_only.status_code == 200, profile_only.text
    assert profile_only.json()["resumeVersion"] is None
    assert profile_only.json()["profile"]["personal"]["name"] == "示例用户"
    assert client.get("/api/v1/resume-versions").json()["total"] == 1


def test_confirm_both_is_atomic_and_profile_merge_is_selected_only(
    client: TestClient,
) -> None:
    existing = client.put(
        "/api/v1/autofill-profile",
        headers=WEB_HEADERS,
        json={
            "personal": {
                "name": "当前姓名",
                "phone": "13000000000",
                "email": "current@example.invalid",
                "currentCity": "北京",
            },
            "education": [{"school": "既有大学", "major": "数学"}],
            "experience": [{"company": "既有公司", "position": "助理"}],
            "links": {"github": "https://github.com/current"},
        },
    )
    assert existing.status_code == 200

    response = client.post(
        "/api/v1/resume-imports/confirm",
        headers=WEB_HEADERS,
        json={
            "resumeVersion": {"name": "导入简历 2026-09-07", "content": "确认后的正文"},
            "profileImport": {
                "personal": {"phone": "13800138000"},
                "education": [{"school": "新增大学", "major": "信息管理"}],
                "experience": [{"company": "新增公司", "position": "产品实习生"}],
                "links": {},
            },
        },
    )

    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["resumeVersion"]["content"] == "确认后的正文"
    profile = payload["profile"]
    assert profile["personal"] == {
        "name": "当前姓名",
        "phone": "13800138000",
        "email": "current@example.invalid",
        "currentCity": "北京",
    }
    assert [item["school"] for item in profile["education"]] == ["既有大学", "新增大学"]
    assert [item["company"] for item in profile["experience"]] == ["既有公司", "新增公司"]
    assert profile["links"]["github"] == "https://github.com/current"


def test_combined_confirm_rolls_back_resume_when_profile_would_exceed_limit(
    client: TestClient,
) -> None:
    existing_rows = [{"school": f"示例大学 {index}"} for index in range(20)]
    saved = client.put(
        "/api/v1/autofill-profile",
        headers=WEB_HEADERS,
        json={"personal": {}, "education": existing_rows, "experience": [], "links": {}},
    )
    assert saved.status_code == 200

    response = client.post(
        "/api/v1/resume-imports/confirm",
        headers=WEB_HEADERS,
        json={
            "resumeVersion": {"name": "不应保存", "content": "不应保存的正文"},
            "profileImport": {"education": [{"school": "第 21 条"}]},
        },
    )

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "VALIDATION_ERROR"
    assert client.get("/api/v1/resume-versions").json()["total"] == 0
    assert len(client.get("/api/v1/autofill-profile").json()["profile"]["education"]) == 20


def test_confirmed_resume_and_profile_persist_after_restart(database_path: Path) -> None:
    settings = ApiSettings.from_environment({"JOBPILOT_CORS_ORIGINS": WEB_ORIGIN})
    first_app = create_app(settings, database_path=database_path)
    with TestClient(first_app, base_url="http://127.0.0.1") as first_client:
        confirmed = first_client.post(
            "/api/v1/resume-imports/confirm",
            headers=WEB_HEADERS,
            json={
                "resumeVersion": {"name": "重启持久化", "content": "确认后的正文"},
                "profileImport": {"personal": {"name": "示例用户"}},
            },
        )
        assert confirmed.status_code == 200

    restarted = create_app(settings, database_path=database_path)
    with TestClient(restarted, base_url="http://127.0.0.1") as restarted_client:
        assert restarted_client.get("/api/v1/resume-versions").json()["total"] == 1
        assert (
            restarted_client.get("/api/v1/autofill-profile").json()["profile"]["personal"]["name"]
            == "示例用户"
        )


def test_confirm_rejects_empty_or_non_json_command(client: TestClient) -> None:
    empty = client.post("/api/v1/resume-imports/confirm", headers=WEB_HEADERS, json={})
    empty_profile_import = client.post(
        "/api/v1/resume-imports/confirm",
        headers=WEB_HEADERS,
        json={"profileImport": {}},
    )
    non_json = client.post(
        "/api/v1/resume-imports/confirm",
        headers={**WEB_HEADERS, "Content-Type": "text/plain"},
        content="{}",
    )

    assert empty.status_code == 422
    assert empty.json()["error"]["code"] == "VALIDATION_ERROR"
    assert empty_profile_import.status_code == 422
    assert empty_profile_import.json()["error"]["code"] == "VALIDATION_ERROR"
    assert non_json.status_code == 415
    assert non_json.json()["error"]["code"] == "JSON_REQUIRED"

    selected_null = client.post(
        "/api/v1/resume-imports/confirm",
        headers=WEB_HEADERS,
        json={"profileImport": {"personal": {"phone": None}}},
    )
    assert selected_null.status_code == 422
    assert selected_null.json()["error"]["code"] == "VALIDATION_ERROR"


def _parse_docx(client: TestClient, content: bytes):
    return client.post(
        "/api/v1/resume-imports/parse",
        headers=WEB_HEADERS,
        files={"file": ("resume.docx", content, DOCX_MIME)},
    )


def _resume_docx(extra: str | None = None) -> bytes:
    document = Document()
    document.add_paragraph("基本信息")
    document.add_paragraph("示例用户")
    document.add_paragraph("电话：13800138000")
    document.add_paragraph("邮箱：candidate@example.invalid")
    document.add_paragraph("教育经历")
    table = document.add_table(rows=4, cols=1)
    table.cell(0, 0).text = "学校：示例大学"
    table.cell(1, 0).text = "专业：信息管理"
    table.cell(2, 0).text = "学历：本科"
    table.cell(3, 0).text = "时间：2022.09 - 2026.06"
    if extra:
        document.add_paragraph(extra)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
