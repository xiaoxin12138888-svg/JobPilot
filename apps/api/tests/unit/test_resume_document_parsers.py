from __future__ import annotations

from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
)

from jobpilot_api.domain.resume_imports import DocumentBlockKind, ResumeImportError
from jobpilot_api.infrastructure.resume_import.docx_parser import parse_docx
from jobpilot_api.infrastructure.resume_import.pdf_parser import parse_pdf


def test_docx_parser_preserves_paragraph_table_and_paragraph_order() -> None:
    document = Document()
    document.add_paragraph("教育经历")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "学校：示例大学"
    table.cell(0, 1).text = "学历：本科"
    table.cell(1, 0).text = "专业：信息管理"
    table.cell(1, 1).text = "2022.09 - 2026.06"
    document.add_paragraph("专业技能\nPython、SQL")

    parsed = parse_docx(_save_docx(document))

    assert [(block.kind, block.text) for block in parsed.blocks] == [
        (DocumentBlockKind.TEXT, "教育经历"),
        (DocumentBlockKind.TABLE_ROW, "学校：示例大学 | 学历：本科"),
        (DocumentBlockKind.TABLE_ROW, "专业：信息管理 | 2022.09 - 2026.06"),
        (DocumentBlockKind.TEXT, "专业技能"),
        (DocumentBlockKind.TEXT, "Python、SQL"),
    ]
    assert parsed.raw_text.splitlines()[1] == "学校：示例大学 | 学历：本科"
    assert "TABLE_ORDER_REVIEW" in [warning.code for warning in parsed.warnings]


def test_docx_parser_rejects_path_traversal_and_zip_bomb_like_entry() -> None:
    traversal = BytesIO()
    with ZipFile(traversal, "w", ZIP_DEFLATED) as archive:
        archive.writestr("../outside.xml", "safe")

    with pytest.raises(ResumeImportError) as traversal_error:
        parse_docx(traversal.getvalue())
    assert traversal_error.value.code == "RESUME_DOCX_INVALID"

    compressed = BytesIO()
    with ZipFile(compressed, "w", ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "A" * 1_000_000)

    with pytest.raises(ResumeImportError) as bomb_error:
        parse_docx(compressed.getvalue())
    assert bomb_error.value.code == "RESUME_DOCX_INVALID"


def test_docx_parser_rejects_dtd_and_does_not_fetch_external_relationships() -> None:
    safe_document = Document()
    safe_document.add_paragraph("Safe resume text")
    safe = _save_docx(safe_document)
    entries = _zip_entries(safe)
    entries["word/document.xml"] = (
        b'<!DOCTYPE x [<!ENTITY secret SYSTEM "file:///not-read">]>' + entries["word/document.xml"]
    )
    with pytest.raises(ResumeImportError) as dtd_error:
        parse_docx(_make_zip(entries))
    assert dtd_error.value.code == "RESUME_DOCX_INVALID"

    entries = _zip_entries(safe)
    relationships = entries["word/_rels/document.xml.rels"].replace(
        b"</Relationships>",
        b'<Relationship Id="remote" Type="urn:test" Target="https://example.invalid/x" '
        b'TargetMode="External"/></Relationships>',
    )
    entries["word/_rels/document.xml.rels"] = relationships

    parsed = parse_docx(_make_zip(entries))

    assert "已忽略 DOCX 外部关系" in [warning.message for warning in parsed.warnings]


def test_docx_parser_rejects_macro_content() -> None:
    document = Document()
    document.add_paragraph("Safe resume text")
    entries = _zip_entries(_save_docx(document))
    entries["word/vbaProject.bin"] = b"not executed"

    with pytest.raises(ResumeImportError) as macro_error:
        parse_docx(_make_zip(entries))

    assert macro_error.value.code == "RESUME_DOCX_INVALID"


def test_pdf_parser_extracts_multiple_pages_and_line_breaks() -> None:
    parsed = parse_pdf(
        _text_pdf(
            [
                ["Education", "School: Example University"],
                ["Experience", "Company: Example Corp", "Email: user@example.invalid"],
            ]
        )
    )

    assert parsed.page_count == 2
    assert "Education\nSchool: Example University" in parsed.raw_text
    assert "Experience\nCompany: Example Corp" in parsed.raw_text
    assert [block.text for block in parsed.blocks][-1] == "Email: user@example.invalid"


def test_pdf_parser_rejects_no_text_and_encrypted_files() -> None:
    blank = PdfWriter()
    blank.add_blank_page(width=612, height=792)
    blank_bytes = BytesIO()
    blank.write(blank_bytes)

    with pytest.raises(ResumeImportError) as no_text_error:
        parse_pdf(blank_bytes.getvalue())
    assert no_text_error.value.code == "RESUME_PDF_NO_TEXT"

    encrypted = PdfWriter()
    encrypted.add_blank_page(width=612, height=792)
    encrypted.encrypt("password")
    encrypted_bytes = BytesIO()
    encrypted.write(encrypted_bytes)

    with pytest.raises(ResumeImportError) as encrypted_error:
        parse_pdf(encrypted_bytes.getvalue())
    assert encrypted_error.value.code == "RESUME_PDF_ENCRYPTED"


def test_pdf_parser_enforces_page_and_parse_time_limits(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    many_pages = PdfWriter()
    for _ in range(101):
        many_pages.add_blank_page(width=612, height=792)
    many_pages_bytes = BytesIO()
    many_pages.write(many_pages_bytes)

    with pytest.raises(ResumeImportError) as page_error:
        parse_pdf(many_pages_bytes.getvalue())
    assert page_error.value.code == "RESUME_PDF_INVALID"

    ticks = iter((0.0, 11.0))
    monkeypatch.setattr(
        "jobpilot_api.infrastructure.resume_import.pdf_parser.monotonic", lambda: next(ticks)
    )
    with pytest.raises(ResumeImportError) as timeout_error:
        parse_pdf(_text_pdf([["Education", "School: Example University"]]))
    assert timeout_error.value.code == "RESUME_PARSE_TIMEOUT"


def _save_docx(document: Document) -> bytes:
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _zip_entries(value: bytes) -> dict[str, bytes]:
    with ZipFile(BytesIO(value)) as archive:
        return {item.filename: archive.read(item) for item in archive.infolist()}


def _make_zip(entries: dict[str, bytes]) -> bytes:
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for name, value in entries.items():
            archive.writestr(name, value)
    return output.getvalue()


def _text_pdf(pages: list[list[str]]) -> bytes:
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    for lines in pages:
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
        )
        commands = ["BT", "/F1 12 Tf", "72 720 Td", "14 TL"]
        for index, line in enumerate(lines):
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            if index:
                commands.append("T*")
            commands.append(f"({escaped}) Tj")
        commands.append("ET")
        stream = DecodedStreamObject()
        stream.set_data("\n".join(commands).encode("latin-1"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()
